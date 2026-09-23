from __future__ import annotations

from io import BytesIO

from pptx import Presentation


def load_pptx(
    file_bytes: bytes,
    file_name: str,
) -> list[dict]:

    presentation = Presentation(
        BytesIO(file_bytes)
    )

    documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        texts = []

        for shape in slide.shapes:

            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                texts.append(text)

        if not texts:
            continue

        documents.append(
            {
                "text": "\n".join(texts),
                "source_file": file_name,
                "page_number": slide_number,
            }
        )

    return documents